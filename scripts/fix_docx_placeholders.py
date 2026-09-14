import docx
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

course_dir = r"d:\HAVEN\Ứng dụng trí tuệ nhân tạo KHMT K23A-20260805T083400Z-1-001"

# 1. Update 01_GenAI_SoftwareDevelopment_project-plan.docx
doc1 = docx.Document(os.path.join(course_dir, "01_GenAI_SoftwareDevelopment_project-plan.docx"))
doc1.paragraphs[1].text = "Nhóm 04 - Nhóm sinh viên thực hiện (Lớp KHMT K23A)"
doc1.paragraphs[2].text = "1. Vũ Ngọc Sơn (Trưởng nhóm, Kiến trúc hệ thống & Fullstack)"
doc1.paragraphs[3].text = "2. Vũ Bảo Linh (Kỹ sư Dữ liệu & Tích hợp AI Engine)  |  3. Tô Văn Quyền (Kiểm thử Phần mềm & Đặc tả Nghiệp vụ)  |  4. Lê Bình Nguyên (Thiết kế Giao diện UI/UX & Tài liệu Học thuật)"
doc1.save(os.path.join(course_dir, "01_GenAI_SoftwareDevelopment_project-plan.docx"))
print("✅ Updated 01_GenAI_SoftwareDevelopment_project-plan.docx")

# 2. Update 02_GenAI_SoftwareDevelopment_requirements-qa.docx
doc2 = docx.Document(os.path.join(course_dir, "02_GenAI_SoftwareDevelopment_requirements-qa.docx"))
doc2.paragraphs[1].text = "Nhóm 04 - Nhóm sinh viên thực hiện (Lớp KHMT K23A)"
doc2.paragraphs[2].text = "1. Vũ Ngọc Sơn (Trưởng nhóm, Kiến trúc hệ thống & Fullstack)"
doc2.paragraphs[3].text = "2. Vũ Bảo Linh (Kỹ sư Dữ liệu & Tích hợp AI Engine)  |  3. Tô Văn Quyền (Kiểm thử Phần mềm & Đặc tả Nghiệp vụ)  |  4. Lê Bình Nguyên (Thiết kế Giao diện UI/UX & Tài liệu Học thuật)"
doc2.save(os.path.join(course_dir, "02_GenAI_SoftwareDevelopment_requirements-qa.docx"))
print("✅ Updated 02_GenAI_SoftwareDevelopment_requirements-qa.docx")

# 3. Update 03_GenAI_SoftwareDevelopment_requirements-specification.docx
doc3 = docx.Document(os.path.join(course_dir, "03_GenAI_SoftwareDevelopment_requirements-specification.docx"))
doc3.paragraphs[0].text = "TÀI LIỆU ĐẶC TẢ YÊU CẦU PHẦN MỀM (SRS)"
doc3.paragraphs[1].text = "Tên ứng dụng: Hệ thống Nền tảng PropTech Cho Thuê & Quản Lý Căn Hộ Tích Hợp Trí Tuệ Nhân Tạo (HAVEN)"
doc3.paragraphs[2].text = "Mã dự án: HAVEN-PROPTECH-K23A  |  Môn học: Ứng dụng Trí tuệ Nhân tạo trong Phát triển Phần mềm (KHMT K23A)"
doc3.paragraphs[3].text = "Nhóm thực hiện: 1. Vũ Ngọc Sơn (Trưởng nhóm)  |  2. Vũ Bảo Linh  |  3. Tô Văn Quyền  |  4. Lê Bình Nguyên"

doc3.paragraphs[6].text = "Tài liệu Đặc tả Yêu cầu Phần mềm (Software Requirements Specification - SRS) này mô tả toàn diện, chi tiết và có hệ thống các yêu cầu chức năng, phi chức năng, các ràng buộc kỹ thuật và mô hình thiết kế của hệ thống HAVEN PropTech Platform. Tài liệu đóng vai trò là cơ sở kỹ thuật thống nhất cho các thành viên phát triển, kiểm thử viên và giảng viên đánh giá môn học."

doc3.paragraphs[9].text = "Nền tảng PropTech web ứng dụng đa phân hệ dành cho thị trường thuê và cho thuê căn hộ tại các đô thị lớn tại Việt Nam (16 thành phố: Hà Nội, TP. Hồ Chí Minh, Đà Nẵng, Hải Phòng, Bình Dương, Nha Trang, Cần Thơ, Vũng Tàu, Hạ Long, Đà Lạt, Huế, Quy Nhơn, Biên Hòa, Vinh, Thanh Hóa, Buôn Ma Thuột). Hệ thống bao gồm 4 phân hệ con: (1) Phân hệ Người Thuê (Tenant Experience); (2) Phân hệ Chủ Nhà & Vận Hành (Landlord Operations); (3) Phân hệ Quản Trị Sàn (Marketplace Admin); (4) Phân hệ Cư Dân (Resident App)."

doc3.paragraphs[10].text = "Tài liệu này phục vụ cho: Đội ngũ kỹ sư phát triển phần mềm (Frontend, Backend, AI Engine), Kỹ sư kiểm thử chất lượng (QA/QC), Nhà quản lý sản phẩm (Product Owner), và Giảng viên hướng dẫn / Hội đồng đánh giá môn học."

doc3.paragraphs[13].text = "Bảng định nghĩa các thuật ngữ chuyên ngành bất động sản, công nghệ PropTech và các từ viết tắt được sử dụng xuyên suốt trong toàn bộ tài liệu đặc tả hệ thống HAVEN (Chi tiết xem Bảng 1 bên dưới)."

doc3.paragraphs[19].text = "Mô hình Use Case tổng quát của hệ thống HAVEN thể hiện mối quan hệ tương tác giữa 5 Tác nhân chính (Người Thuê, Chủ Nhà, Cư Dân, Quản Trị Viên, Dịch Vụ GenAI) với 12 Use Case cốt lõi: UC001 Tìm kiếm True Cost & Bản Đồ, UC002 AI Housing Advisor, UC003 So Sánh Radar, UC004 Chat Smart Chips & Đặt Lịch, UC005 Ký Hợp Đồng & Cọc Escrow, UC006 Bàn Giao 15 Mục, UC007 Cẩm Nang & Commute, UC008 Virtual Tour 360, UC009 Smart Listing AI, UC010 Quản Lý Căn Hộ & Thu Tiền, UC011 Quản Trị Sàn & Duyệt Tin, UC012 IoT Smart Home & Yêu Cầu Sửa Chữa."

doc3.paragraphs[22].text = "Hệ thống HAVEN phân định rõ quyền hạn và vai trò của 5 tác nhân chính tham gia tương tác với nền tảng (Chi tiết xem Bảng 3 bên dưới)."

doc3.paragraphs[25].text = "Danh sách 12 Use Case chính tương ứng với các yêu cầu chức năng (FR-01 đến FR-15) được thu thập từ tài liệu 02_requirements-qa.docx, phân loại theo độ ưu tiên nghiệp vụ P0/P1/P2 (Chi tiết xem Bảng 4 bên dưới)."

doc3.paragraphs[27].text = "Quy cách đánh số: UC: Use Case, 001, 002... là số thứ tự; FR: Functional Requirement; NFR: Non-Functional Requirement."

doc3.paragraphs[30].text = "Hệ thống được xây dựng trên nền tảng Web hiện đại: Ngôn ngữ TypeScript / JavaScript; Thư viện React 19, Vite 8, Tailwind CSS v4, Lucide Icons, Three.js / WebGL shader cho giao diện trực quan; Trí tuệ nhân tạo tích hợp Google Gemini API (NLP Semantic Parsing & Computer Vision); Môi trường vận hành tương thích trên mọi trình duyệt hiện đại (Chrome, Edge, Safari, Firefox) và thiết bị (Desktop, Tablet, Mobile Responsive)."

doc3.paragraphs[33].text = "Phần này đặc tả chi tiết 2 Use Case tiêu biểu cốt lõi đại diện cho Phân hệ Người Thuê (UC001) và Phân hệ Chủ Nhà kết hợp Trí Tuệ Nhân Tạo (UC009). Cả hai Use Case đều được chuẩn hóa theo quy chuẩn IEEE 830 với đầy đủ Điều kiện tiên quyết, Luồng sự kiện chính, Luồng ngoại lệ và Kịch bản kiểm thử."

doc3.paragraphs[35].text = "UC001: Tìm Kiếm Căn Hộ Theo Chi Phí Thực Tế (True Cost) & Bản Đồ An Toàn PCCC / Ngập Lụt"
doc3.paragraphs[36].text = "Use case cho phép người thuê tìm kiếm, lọc căn hộ theo ngân sách thực tế (đã gồm điện nước, gửi xe, phí quản lý) và trực quan hóa các lớp rủi ro PCCC, ngập lụt trước khi quyết định thuê."

doc3.paragraphs[39].text = "Luồng tuần tự (Sequence Flow): Tenant -> UI SearchView -> TrueCostEngine (Tính toán chi phí) -> ConfidenceMapService (Lấy lớp rủi ro PCCC QCVN 06 & Triều cường) -> Trả về danh sách căn hộ đạt chuẩn và bản đồ nhiệt trực quan."

doc3.paragraphs[41].text = "UC009: Đăng Tin Tự Động Bằng Trí Tuệ Nhân Tạo (Smart Listing Creator)"
doc3.paragraphs[42].text = "Use case cho phép chủ nhà tải ảnh căn hộ thực tế; hệ thống GenAI tự động phân tích loại phòng, sinh tiêu đề chuẩn SEO, soạn thảo mô tả hấp dẫn và gợi ý khung giá thuê tối ưu trong 60 giây."

doc3.paragraphs[45].text = "Luồng tuần tự (Sequence Flow): Landlord -> Tải 3-8 ảnh phòng -> GenAI Engine (Computer Vision nhận diện không gian & NLP sinh content) -> Đề xuất giá thị trường -> Chủ nhà duyệt & Bấm xuất bản -> Marketplace gắn nhãn Verified Listing."

doc3.paragraphs[48].text = "Hệ thống HAVEN đã được hoàn thiện giai đoạn 1 và giai đoạn 2: Mã nguồn sạch 100% không lỗi TypeScript (0 errors), đã tích hợp sẵn kho dữ liệu 150 căn hộ mẫu tại 16 đô thị lớn tại Việt Nam, hỗ trợ cả 2 chế độ Dark Mode / Light Mode mượt mà. Hệ thống sẵn sàng cho bài kiểm tra tiến độ Buổi 2 và có thể chạy trực tiếp bằng lệnh 'npm run dev' hoặc bản dựng 'npm run build'."

doc3.save(os.path.join(course_dir, "03_GenAI_SoftwareDevelopment_requirements-specification.docx"))
print("✅ Updated 03_GenAI_SoftwareDevelopment_requirements-specification.docx")
